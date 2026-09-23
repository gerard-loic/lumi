import io
import os

from lib.pipelines.block import Block
from lib.pipelines.pipelinecontext import PipelineContext
from lib.pipelines.utils.filesource import resolve_file_source, FileSourceError
from lib.pipelines.utils.tabular import build_records, to_output
from lib.log.logger import Logger, ERROR, OK

#Bloc ExcelReader : lit un classeur Excel (.xlsx / .xlsm via openpyxl, .xls via xlrd) et écrit dans
#le contexte une table exploitable par le bloc DataView (list[dict], ou forme colonnaire). Le fichier
#peut être produit pendant le run (sortie d'un bloc DataViewFile / Agent, URL "/files/...", clé
#FileStore) ou situé ailleurs sur le serveur (chemin).
#
#Paramètres de configuration (clé "config" du bloc) :
#  - source      (str|dict, obligatoire) : fichier à lire. Voir lib/pipelines/utils/filesource.py.
#  - format      (str,  défaut auto)     : "xlsx" | "xlsm" | "xls" ; sinon déduit de l'extension du
#                                          nom de fichier, avec repli sur "xlsx".
#  - sheet       (str|int|list, défaut None) : feuille(s) à lire. Nom, index 0-based, ou "*" / liste
#                                          -> le résultat est alors {nom_feuille: table}. None -> 1re feuille.
#  - has_header  (bool, défaut True)     : la première ligne porte les noms de colonnes.
#  - columns     (list, défaut [])       : noms de colonnes imposés (si has_header, la 1re ligne est sautée ;
#                                          sinon et si vide, colonnes nommées "col_1", "col_2", ...).
#  - skip_rows   (int,  défaut 0)        : lignes ignorées en tête de feuille (avant l'en-tête).
#  - skip_empty  (bool, défaut True)     : ignore les lignes entièrement vides.
#  - trim        (bool, défaut True)     : retire les espaces de début / fin des cellules chaine.
#  - limit       (int,  défaut 0)        : nombre maximum de lignes de données par feuille (0 = pas de limite).
#  - as          (str,  défaut "list")   : "list" -> list[dict] ; "dict" -> {colonne: [valeurs]}.
#  - output      (str,  défaut "result") : clé du contexte où stocker le résultat.
#
#Les dates / heures des cellules sont normalisées en chaine ISO 8601 (sérialisables en JSON).
class ExcelReader(Block):
    def __init__(self, block_uid:str, config:dict, on_success_block:str=None, on_error_block:str=None):
        super().__init__("ExcelReader", block_uid, config, on_success_block=on_success_block, on_error_block=on_error_block)

    def execute(self, context:PipelineContext):
        fmt        = str(context.getConfig(key="format", default="")).lower().lstrip(".")
        sheet      = context.getConfig(key="sheet", default=None)
        has_header = bool(context.getConfig(key="has_header", default=True))
        columns    = context.getConfig(key="columns", default=[]) or []
        skip_rows  = int(context.getConfig(key="skip_rows", default=0) or 0)
        skip_empty = bool(context.getConfig(key="skip_empty", default=True))
        trim       = bool(context.getConfig(key="trim", default=True))
        limit      = int(context.getConfig(key="limit", default=0) or 0)
        as_        = str(context.getConfig(key="as", default="list")).lower()
        output     = context.getConfig(key="output", default="result")

        try:
            content, filename = resolve_file_source(self._config.get("source"), context)
        except FileSourceError as e:
            Logger.write(f"[Block ExcelReader] {e}", type=ERROR)
            return False

        kind = fmt or os.path.splitext(filename or "")[1].lower().lstrip(".") or "xlsx"

        try:
            sheets = self._read_xls(content) if kind == "xls" else self._read_xlsx(content)
        except Exception as e:
            Logger.write(f"[Block ExcelReader] Cannot read workbook ({kind}) : {e}", type=ERROR)
            return False

        if not sheets:
            Logger.write("[Block ExcelReader] Workbook has no sheet", type=ERROR)
            return False

        #Sélection multi-feuilles : "*" ou une liste -> résultat {nom_feuille: table}.
        multi = sheet == "*" or isinstance(sheet, (list, tuple))
        if multi:
            wanted = list(sheets) if sheet == "*" else [str(s) for s in sheet]
        else:
            wanted = [self._pick_sheet(sheets, sheet)]

        tables = {}
        total = 0
        for name in wanted:
            if name not in sheets:
                Logger.write(f"[Block ExcelReader] Sheet '{name}' not found", type=ERROR)
                return False
            records = build_records(
                sheets[name],
                has_header=has_header,
                columns=list(columns) or None,
                skip_rows=skip_rows,
                skip_empty=skip_empty,
                trim=trim,
                limit=limit,
            )
            tables[name] = to_output(records, as_)
            total += len(records)

        result = tables if multi else tables[wanted[0]]
        context.set(output, result)

        scope = f"{len(wanted)} sheet(s)" if multi else f"sheet '{wanted[0]}'"
        Logger.write(f"[Block ExcelReader] {filename} {scope} -> '{output}' ({total} row(s))", type=OK)
        return True

    #Résout la feuille demandée (None -> première, int -> index 0-based, str -> nom) en son nom.
    @staticmethod
    def _pick_sheet(sheets:dict, sheet):
        names = list(sheets)
        if sheet is None or sheet == "":
            return names[0]
        if isinstance(sheet, bool):
            return names[0]
        if isinstance(sheet, int):
            return names[sheet] if -len(names) <= sheet < len(names) else names[0]
        text = str(sheet)
        if text in sheets:
            return text
        if text.lstrip("-").isdigit():
            idx = int(text)
            if -len(names) <= idx < len(names):
                return names[idx]
        return text  #laissé tel quel -> déclenchera l'erreur "Sheet not found" en amont

    #{nom_feuille: matrice (liste de listes)} pour un classeur .xlsx / .xlsm.
    @staticmethod
    def _read_xlsx(content:bytes)->dict:
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            return {ws.title: [list(row) for row in ws.iter_rows(values_only=True)]
                    for ws in wb.worksheets}
        finally:
            wb.close()

    #{nom_feuille: matrice (liste de listes)} pour un classeur .xls (format binaire hérité).
    @staticmethod
    def _read_xls(content:bytes)->dict:
        import xlrd

        book = xlrd.open_workbook(file_contents=content)
        sheets = {}
        for sh in book.sheets():
            matrix = []
            for rx in range(sh.nrows):
                row = []
                for cx in range(sh.ncols):
                    cell = sh.cell(rx, cx)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode).isoformat()
                    elif cell.ctype == xlrd.XL_CELL_EMPTY:
                        value = None
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    row.append(value)
                matrix.append(row)
            sheets[sh.name] = matrix
        return sheets
